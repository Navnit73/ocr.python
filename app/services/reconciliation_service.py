"""
Bank Reconciliation Orchestration Service.
Fans a submission out to the existing OCR job pipeline (PDF/images) or the spreadsheet parser (CSV/XLSX),
runs the deterministic reconciliation engine once both extractions finish, persists the report in MongoDB,
and dispatches signed, SSRF-protected completion/failure webhooks without blocking workers.
"""

import asyncio
import csv
import hashlib
import io
import ipaddress
import logging
import os
import secrets
import socket
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from typing import Any, Dict, List, Optional, Set, Tuple
from urllib.parse import urlparse

import pymupdf
from fastapi import HTTPException, UploadFile, status
from pymongo.errors import DocumentTooLarge

from app.core.config import get_settings
from app.db.repositories.job_repo import JobRepository
from app.db.repositories.reconciliation_repo import ReconciliationRepository
from app.services.export_service import sanitize_for_formula_injection
from app.services.file_validator import FileValidator
from app.services.pdf_service import PDFService
from app.services.reconciliation_engine import (
    ReconciliationError,
    ReconciliationOptions,
    parse_spreadsheet,
    reconcile,
)
from app.services.storage_service import StorageService
from app.services.webhook_service import WebhookService

logger = logging.getLogger("reconciliation")

SPREADSHEET_EXTENSIONS = (".csv", ".xlsx")
ROLES = ("bank_statement", "business_ledger")
ROLE_LABELS = {"bank_statement": "bank statement", "business_ledger": "business ledger"}

# Fixed upload error texts (after the "<Label>: " prefix) that clients match on to ask for a password.
PDF_PASSWORD_REQUIRED = "This PDF is password-protected. Provide its password to continue."
PDF_PASSWORD_INCORRECT = "Incorrect PDF password."
TERMINAL_JOB_STATUSES = ("completed", "failed", "cancelled")
WEBHOOK_MAX_ATTEMPTS = 4  # backoff 1s, 2s, 4s between attempts (bounded)
STALE_CLAIM_SECONDS = 600

RESULT_SECTIONS = {
    "matched": "matched_transactions",
    "unmatched_bank": "unmatched_bank_transactions",
    "unmatched_ledger": "unmatched_ledger_transactions",
    "possible_duplicates": "possible_duplicates",
    "amount_mismatches": "amount_mismatches",
}

# Keep strong references so fire-and-forget tasks are not garbage collected mid-flight.
_background_tasks: Set[asyncio.Task] = set()


def _spawn(coro) -> None:
    task = asyncio.create_task(coro)
    _background_tasks.add(task)
    task.add_done_callback(_background_tasks.discard)


def hash_access_token(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def as_utc(value: Optional[datetime]) -> Optional[datetime]:
    """Mongo drivers may return naive UTC datetimes; normalize for comparisons and serialization."""
    if value is None or not isinstance(value, datetime):
        return value
    return value if value.tzinfo else value.replace(tzinfo=timezone.utc)


class ReconciliationService:
    """Coordinates reconciliation submissions end to end."""

    # ------------------------------------------------------------------
    # Validation
    # ------------------------------------------------------------------

    @staticmethod
    def unlock_pdf(content: bytes, password: Optional[str]) -> bytes:
        """
        Returns the PDF unchanged when it opens without a password, or an unencrypted copy when it
        needs one and `password` is correct. Only the unlocked copy is stored and processed, so the
        password itself is never persisted. Raises ValueError(PDF_PASSWORD_REQUIRED / _INCORRECT).
        """
        try:
            doc = pymupdf.open(stream=content, filetype="pdf")
        except Exception:
            return content  # Damaged files are reported by the normal PDF validation.
        try:
            if not doc.needs_pass:
                return content
            if not password:
                raise ValueError(PDF_PASSWORD_REQUIRED)
            if not doc.authenticate(password):
                raise ValueError(PDF_PASSWORD_INCORRECT)
            return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_NONE)
        finally:
            doc.close()

    @classmethod
    async def validate_document(cls, upload: UploadFile, role: str, password: Optional[str] = None) -> Tuple[str, bytes]:
        """
        Returns (kind, bytes) where kind is 'pdf', 'image' or 'spreadsheet'. A password-protected PDF
        is unlocked with `password` and its unencrypted bytes are returned.
        """
        settings = get_settings()
        label = ROLE_LABELS[role]
        filename = upload.filename or ""
        ext = os.path.splitext(filename.lower())[1]
        if ext not in SPREADSHEET_EXTENSIONS:
            if ext not in settings.allowed_extensions:
                allowed = ", ".join(list(settings.allowed_extensions) + list(SPREADSHEET_EXTENSIONS))
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Unsupported {label} file type '{ext or 'none'}'. Allowed: {allowed}",
                )
            try:
                kind, content = await FileValidator.validate_upload(upload)
                if kind == "pdf":
                    try:
                        content = await asyncio.to_thread(cls.unlock_pdf, content, password)
                    except ValueError as e:
                        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
                    # Rejects PDFs above MAX_PDF_PAGES (and any still locked) before any job is queued.
                    await PDFService.get_pdf_page_count_async(content)
            except HTTPException as e:
                detail = str(e.detail)
                if "password" in detail.lower() and detail not in (PDF_PASSWORD_REQUIRED, PDF_PASSWORD_INCORRECT):
                    detail = PDF_PASSWORD_REQUIRED
                raise HTTPException(status_code=e.status_code, detail=f"{label.capitalize()}: {detail}")
            return kind, content

        content = await upload.read()
        try:
            FileValidator.validate_bytes_and_size(content, filename)
        except HTTPException as e:
            raise HTTPException(status_code=e.status_code, detail=f"{label.capitalize()}: {e.detail}")
        if ext == ".xlsx" and not content.startswith(b"PK\x03\x04"):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"The {label} is not a valid .xlsx file.")
        if ext == ".csv" and (b"\x00" in content[:8192] or content.startswith((b"%PDF", b"PK\x03\x04"))):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"The {label} is not a valid text CSV file.")
        return "spreadsheet", content

    @classmethod
    def validate_webhook_url(cls, url: str) -> str:
        """
        Allows only HTTPS callbacks whose host resolves exclusively to public IP addresses.
        Blocks credentials-in-URL, private/loopback/link-local/reserved/multicast targets.
        Redirects are never followed at delivery time.
        """
        url = (url or "").strip()
        parsed = urlparse(url)
        if parsed.scheme != "https" or not parsed.hostname:
            raise ValueError("webhook_url must be an absolute HTTPS URL.")
        if parsed.username or parsed.password:
            raise ValueError("webhook_url must not contain credentials.")
        if len(url) > 2048:
            raise ValueError("webhook_url is too long.")
        host = parsed.hostname
        if host.lower() in ("localhost",) or host.lower().endswith((".localhost", ".internal", ".local")):
            raise ValueError("webhook_url must point to a public host.")
        try:
            infos = socket.getaddrinfo(host, parsed.port or 443, proto=socket.IPPROTO_TCP)
        except socket.gaierror:
            raise ValueError("webhook_url host could not be resolved.")
        for info in infos:
            ip = ipaddress.ip_address(info[4][0].split("%")[0])
            if ip.version == 6 and ip.ipv4_mapped:
                ip = ip.ipv4_mapped
            if (
                ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved
                or ip.is_multicast or ip.is_unspecified or not ip.is_global
            ):
                raise ValueError("webhook_url resolves to a private or internal address.")
        return url

    # ------------------------------------------------------------------
    # Submission
    # ------------------------------------------------------------------

    @classmethod
    async def create_reconciliation(
        cls,
        documents: Dict[str, Tuple[str, str, bytes, Optional[str]]],
        options: ReconciliationOptions,
        owner_id: Optional[str],
        owner_email: Optional[str],
        webhook_url: Optional[str],
        webhook_secret: Optional[str],
    ) -> Tuple[Dict[str, Any], Optional[str]]:
        """
        documents: role -> (kind, filename, bytes, content_type).
        Returns (record, guest_access_token). The token is only generated for guests.
        """
        settings = get_settings()
        now = datetime.now(timezone.utc)
        reconciliation_id = f"rec_{uuid.uuid4().hex[:16]}"
        access_token = None if owner_id else secrets.token_urlsafe(32)
        job_owner = owner_id or f"guest:{reconciliation_id}"
        user_email = (owner_email or "guest").lower().strip() if owner_id else "guest"

        doc_entries: Dict[str, Dict[str, Any]] = {}
        jobs_to_submit: List[str] = []
        record = {
            "reconciliation_id": reconciliation_id,
            "owner_id": owner_id,
            "owner_email": user_email if owner_id else None,
            "access_token_hash": hash_access_token(access_token) if access_token else None,
            "expires_at": None if owner_id else now + timedelta(days=settings.storage_retention_days),
            "documents": doc_entries,
            "status": "queued",
            "stage": "extracting",
            "progress": 0,
            "message": "Documents uploaded. Reconciliation has been queued.",
            "options": {
                "date_tolerance_days": options.date_tolerance_days,
                "description_matching": options.description_matching,
            },
            "result": None,
            "error": None,
            "webhook": {
                "url": webhook_url,
                "secret": webhook_secret or None,
                "delivery_status": "pending" if webhook_url else "not_configured",
            },
            "created_at": now,
            "updated_at": now,
            "completed_at": None,
        }
        try:
            await cls._prepare_documents(
                reconciliation_id, documents, job_owner, user_email, doc_entries, jobs_to_submit
            )
            # Persist before submitting jobs so a fast worker always finds the record.
            await ReconciliationRepository.create(record)
        except Exception:
            if jobs_to_submit:
                await JobRepository._collection().delete_many({"job_id": {"$in": jobs_to_submit}})
            for entry in doc_entries.values():
                await StorageService.delete_document_directory(entry["document_id"])
            raise

        from app.workers.worker_manager import WorkerManager

        for job_id in jobs_to_submit:
            await WorkerManager.submit_job(job_id, user_email=user_email)
        if not jobs_to_submit:
            _spawn(cls.try_finalize(reconciliation_id))
        return record, access_token

    @classmethod
    async def _prepare_documents(
        cls,
        reconciliation_id: str,
        documents: Dict[str, Tuple[str, str, bytes, Optional[str]]],
        job_owner: str,
        user_email: str,
        doc_entries: Dict[str, Dict[str, Any]],
        jobs_to_submit: List[str],
    ) -> None:
        """Saves both uploads and creates (but does not submit) OCR jobs for PDF/image documents."""
        for role in ROLES:
            kind, filename, content, content_type = documents[role]
            document_id = f"{reconciliation_id}_{'bank' if role == 'bank_statement' else 'ledger'}"
            file_ref, checksum, size = await StorageService.save_file(content, filename, document_id)
            entry: Dict[str, Any] = {
                "document_id": document_id,
                "filename": filename,
                "kind": kind,
                "file_size_bytes": size,
                "checksum_sha256": checksum,
                "job_id": None,
                "file_reference": file_ref,
            }
            if kind != "spreadsheet":
                job_id = f"job_{uuid.uuid4().hex[:12]}"
                entry["job_id"] = job_id
                await JobRepository.create_job({
                    "job_id": job_id,
                    "document_id": document_id,
                    "user_id": job_owner,
                    "user_email": user_email,
                    "status": "queued",
                    "total_pages": 1,
                    "current_stage": "queued",
                    "message": f"Queued for reconciliation {reconciliation_id}",
                    "file_reference": file_ref,
                    "document_type": "bank_statement",
                    "language": "en",
                    "clean_with_ai": True,
                    "metadata": {
                        "filename": filename,
                        "content_type": content_type,
                        "file_size_bytes": size,
                        "checksum_sha256": checksum,
                        "doc_category": kind,
                        "user_email": user_email,
                        "reconciliation_id": reconciliation_id,
                        "reconciliation_role": role,
                    },
                })
                jobs_to_submit.append(job_id)
            doc_entries[role] = entry

    # ------------------------------------------------------------------
    # Completion
    # ------------------------------------------------------------------

    @classmethod
    async def on_extraction_job_finished(cls, job: Dict[str, Any]) -> None:
        """Hook called by the worker after an OCR job reaches a terminal state."""
        reconciliation_id = (job.get("metadata") or {}).get("reconciliation_id")
        if reconciliation_id:
            # A Celery task closes its event loop when it returns, which would kill a fire-and-forget
            # webhook task, so in Celery mode delivery is awaited inside the task instead.
            await cls.try_finalize(reconciliation_id, await_webhook=get_settings().use_celery)

    @classmethod
    async def try_finalize(cls, reconciliation_id: str, await_webhook: bool = False) -> None:
        """Runs reconciliation once both extractions are terminal. Safe to call repeatedly."""
        record = await ReconciliationRepository.get(reconciliation_id)
        if not record or record.get("status") in ("completed", "failed"):
            return
        jobs = await cls._load_jobs(record, include_result=False)
        if any(j.get("status") not in TERMINAL_JOB_STATUSES for j in jobs.values()):
            return
        if not await ReconciliationRepository.claim_for_reconciliation(reconciliation_id, STALE_CLAIM_SECONDS):
            return

        try:
            statements: Dict[str, Dict[str, Any]] = {}
            jobs = await cls._load_jobs(record, include_result=True)
            for role in ROLES:
                statements[role] = await cls._load_statement(record["documents"][role], jobs, role)
            opts = record.get("options") or {}
            options = ReconciliationOptions(
                date_tolerance_days=int(opts.get("date_tolerance_days", 2)),
                description_matching=bool(opts.get("description_matching", True)),
            )
            result = await asyncio.to_thread(
                reconcile, statements["bank_statement"], statements["business_ledger"], options
            )
        except ReconciliationError as e:
            await cls._finish_failed(record, str(e), await_webhook)
            return
        except Exception as e:
            logger.error(f"Reconciliation {reconciliation_id} failed unexpectedly: {type(e).__name__}", exc_info=True)
            await cls._finish_failed(record, "Reconciliation failed due to an internal processing error.", await_webhook)
            return

        try:
            completed = await ReconciliationRepository.complete(reconciliation_id, result)
        except DocumentTooLarge:
            await cls._finish_failed(
                record,
                "The reconciliation report is too large to store. Split the documents into smaller periods and retry.",
                await_webhook,
            )
            return
        except Exception as e:
            # Leave the claim in place; it becomes stale and is retried by the next status poll.
            logger.error(f"Could not save reconciliation {reconciliation_id}: {type(e).__name__}")
            return
        if completed:
            logger.info(f"Reconciliation {reconciliation_id} completed.")
            await cls._apply_guest_retention(record)
            await cls._notify(record, {
                "event": "reconciliation.completed",
                "reconciliation_id": reconciliation_id,
                "status": "completed",
                "result_url": f"/api/v1/reconciliation/{reconciliation_id}",
            }, await_webhook)

    @classmethod
    async def _finish_failed(cls, record: Dict[str, Any], error: str, await_webhook: bool = False) -> None:
        reconciliation_id = record["reconciliation_id"]
        if await ReconciliationRepository.fail(reconciliation_id, error):
            logger.warning(f"Reconciliation {reconciliation_id} failed: {error}")
            await cls._apply_guest_retention(record)
            await cls._notify(record, {
                "event": "reconciliation.failed",
                "reconciliation_id": reconciliation_id,
                "status": "failed",
                "error": error,
            }, await_webhook)

    @classmethod
    async def _apply_guest_retention(cls, record: Dict[str, Any]) -> None:
        """Guest OCR output (jobs, documents, extractions) expires together with the guest reconciliation via TTL."""
        expires_at = record.get("expires_at")
        if record.get("owner_id") or not expires_at:
            return
        docs = record.get("documents", {}).values()
        document_ids = [d["document_id"] for d in docs]
        job_ids = [d["job_id"] for d in docs if d.get("job_id")]
        if not job_ids:
            return
        try:
            from app.db.mongodb import MongoDBManager

            db = MongoDBManager.get_db()
            update = {"$set": {"expires_at": expires_at}}
            await db["jobs"].update_many({"job_id": {"$in": job_ids}}, update)
            await db["documents"].update_many({"document_id": {"$in": document_ids}}, update)
            await db["extractions"].update_many({"id": {"$in": document_ids}}, update)
        except Exception as e:
            logger.warning(f"Could not set guest retention for {record['reconciliation_id']}: {type(e).__name__}")

    @classmethod
    async def _load_jobs(cls, record: Dict[str, Any], include_result: bool) -> Dict[str, Dict[str, Any]]:
        job_ids = [d["job_id"] for d in record.get("documents", {}).values() if d.get("job_id")]
        if not job_ids:
            return {}
        projection = {"_id": 0, "job_id": 1, "status": 1, "progress": 1}
        if include_result:
            projection.update({"result.extraction": 1, "result.status": 1, "result.raw_text": 1})
        cursor = JobRepository._collection().find({"job_id": {"$in": job_ids}}, projection)
        return {j["job_id"]: j for j in await cursor.to_list(length=len(job_ids))}

    @classmethod
    async def _load_statement(cls, doc: Dict[str, Any], jobs: Dict[str, Dict[str, Any]], role: str) -> Dict[str, Any]:
        label = ROLE_LABELS[role]
        if doc["kind"] == "spreadsheet":
            try:
                content = await StorageService.read_file(doc["file_reference"])
                data = await asyncio.to_thread(parse_spreadsheet, content, doc["filename"])
            except ReconciliationError as e:
                raise ReconciliationError(f"Unable to extract transactions from the {label}: {e}")
            except FileNotFoundError:
                raise ReconciliationError(f"The uploaded {label} is no longer available.")
        else:
            job = jobs.get(doc["job_id"]) or {}
            if job.get("status") != "completed":
                raise ReconciliationError(f"Unable to extract transactions from the {label}.")
            result = job.get("result") or {}
            data = dict(result.get("extraction") or {})
            # Printed text lets the engine catch amounts misread with the wrong decimal separator.
            data["_raw_text"] = result.get("raw_text") or ""
        if not data.get("transactions"):
            raise ReconciliationError(f"Unable to extract transactions from the {label}.")
        return data

    # ------------------------------------------------------------------
    # Webhooks
    # ------------------------------------------------------------------

    @classmethod
    async def _notify(cls, record: Dict[str, Any], payload: Dict[str, Any], await_delivery: bool = False) -> None:
        webhook = record.get("webhook") or {}
        if not webhook.get("url"):
            return
        delivery = cls._deliver_webhook(record["reconciliation_id"], webhook["url"], webhook.get("secret"), payload)
        if await_delivery:
            await delivery
        else:
            _spawn(delivery)

    @classmethod
    async def _deliver_webhook(cls, reconciliation_id: str, url: str, secret: Optional[str], payload: Dict[str, Any]) -> None:
        event = payload["event"]
        delivered = False
        try:
            # Re-check at delivery time to narrow DNS-rebinding windows.
            await asyncio.to_thread(cls.validate_webhook_url, url)
            body = dict(payload)
            body["event_id"] = f"evt_{uuid.uuid4().hex[:16]}"
            body["timestamp"] = datetime.now(timezone.utc).isoformat()
            delivered = await WebhookService.send_webhook(
                callback_url=url,
                payload=body,
                secret=secret,
                max_retries=WEBHOOK_MAX_ATTEMPTS,
                event_id=body["event_id"],
                job_id=reconciliation_id,
                event_name=event,
            )
        except ValueError:
            logger.warning(f"Webhook for {reconciliation_id} skipped: destination failed safety validation.")
        except Exception as e:
            logger.warning(f"Webhook for {reconciliation_id} errored: {type(e).__name__}")
        await ReconciliationRepository.record_webhook_delivery(reconciliation_id, event, delivered)

    # ------------------------------------------------------------------
    # Read views
    # ------------------------------------------------------------------

    @classmethod
    async def build_status(cls, record: Dict[str, Any]) -> Dict[str, Any]:
        """Builds the API 2 response from stored state (never reruns OCR or matching)."""
        state = record.get("status", "queued")
        progress = record.get("progress", 0)
        message = record.get("message") or ""

        if state in ("queued", "processing") and record.get("stage") != "reconciling":
            jobs = await cls._load_jobs(record, include_result=False)
            running, total_progress, pending_roles = False, 0, []
            for role in ROLES:
                doc = record["documents"][role]
                if doc["kind"] == "spreadsheet":
                    total_progress += 100
                    continue
                job = jobs.get(doc["job_id"]) or {}
                job_status = job.get("status", "queued")
                total_progress += 100 if job_status in TERMINAL_JOB_STATUSES else int(job.get("progress") or 0)
                if job_status != "queued":
                    running = True
                if job_status not in TERMINAL_JOB_STATUSES:
                    pending_roles.append(role)
            progress = 5 + int(total_progress / len(ROLES) * 0.8)
            if pending_roles:
                if running and state == "queued":
                    state = "processing"
                    await ReconciliationRepository.mark_processing(record["reconciliation_id"])
                if state == "processing":
                    names = " and ".join(ROLE_LABELS[r] for r in pending_roles)
                    message = f"Extracting {names} transactions"
                else:
                    progress = 0
            else:
                # Both extractions are done but no worker has finalized yet (e.g. after a restart).
                state, message = "processing", "Matching transactions and verifying balances"
                _spawn(cls.try_finalize(record["reconciliation_id"]))
        elif state == "processing":
            claimed_at = as_utc(record.get("claimed_at"))
            if claimed_at and datetime.now(timezone.utc) - claimed_at > timedelta(seconds=STALE_CLAIM_SECONDS):
                _spawn(cls.try_finalize(record["reconciliation_id"]))

        return {
            "success": state != "failed",
            "reconciliation_id": record["reconciliation_id"],
            "status": state,
            "progress": min(progress, 100) if state != "completed" else 100,
            "message": message,
            "result": record.get("result") if state == "completed" else None,
            "error": record.get("error") if state == "failed" else None,
            "created_at": as_utc(record.get("created_at")),
            "completed_at": as_utc(record.get("completed_at")),
        }

    @classmethod
    async def get_transactions(
        cls,
        record: Dict[str, Any],
        category: Optional[str],
        page: int,
        page_size: int,
    ) -> Tuple[int, List[Dict[str, Any]]]:
        """Pages through saved result arrays. Without a category, sections are concatenated in report order."""
        summary = (record.get("result") or {}).get("summary") or {}
        categories = [category] if category else list(RESULT_SECTIONS)
        counts = [(c, int(summary.get(c, 0))) for c in categories]
        total = sum(n for _, n in counts)

        skip = (page - 1) * page_size
        remaining = page_size
        items: List[Dict[str, Any]] = []
        for cat, count in counts:
            if remaining <= 0:
                break
            if skip >= count:
                skip -= count
                continue
            take = min(remaining, count - skip)
            items += await ReconciliationRepository.get_result_section(
                record["reconciliation_id"], RESULT_SECTIONS[cat], skip, take
            )
            remaining -= take
            skip = 0
        return total, items

    # ------------------------------------------------------------------
    # Export
    # ------------------------------------------------------------------

    @staticmethod
    def _summary_rows(record: Dict[str, Any]) -> List[List[Any]]:
        result = record["result"]
        summary = result["summary"]
        rows: List[List[Any]] = [
            ["Reconciliation ID", record["reconciliation_id"]],
            ["Bank statement", record["documents"]["bank_statement"]["filename"]],
            ["Business ledger", record["documents"]["business_ledger"]["filename"]],
            ["Created at", str(as_utc(record.get("created_at")) or "")],
            ["Completed at", str(as_utc(record.get("completed_at")) or "")],
        ]
        labels = {
            "currency": "Currency", "bank_transactions": "Bank transactions", "ledger_transactions": "Ledger transactions",
            "amount_corrections": "OCR amount corrections",
            "matched": "Matched", "unmatched_bank": "Unmatched bank", "unmatched_ledger": "Unmatched ledger",
            "possible_duplicates": "Possible duplicate groups", "amount_mismatches": "Amount / direction mismatches",
            "amount_decimal_places": "Amount decimal places",
            "date_tolerance_days": "Date tolerance (days)", "description_matching": "Description matching",
            "ledger_direction_convention": "Ledger debit/credit convention",
        }
        rows += [[label, summary.get(key)] for key, label in labels.items()]
        return rows

    @staticmethod
    def _balance_rows(result: Dict[str, Any]) -> List[List[Any]]:
        bc = result["balance_check"]
        ledger = bc.get("ledger") or {}
        btb = bc.get("book_to_bank") or {}
        rows: List[List[Any]] = [
            ["Metric", "Bank statement", "Business ledger"],
            ["Opening balance", bc.get("bank_opening_balance"), ledger.get("opening_balance")],
            ["Opening balance source", bc.get("opening_balance_source"), ledger.get("opening_balance_source")],
            ["Closing balance", bc.get("bank_closing_balance"), ledger.get("closing_balance")],
            ["Calculated closing balance", bc.get("calculated_closing_balance"), ledger.get("calculated_closing_balance")],
            ["Difference", bc.get("difference"), ledger.get("difference")],
            ["Closing balance check", bc.get("status"), ledger.get("status")],
            ["Running balance check", (bc.get("running_balance") or {}).get("status"), (ledger.get("running_balance") or {}).get("status")],
            ["Running balance rows checked", (bc.get("running_balance") or {}).get("checked"), (ledger.get("running_balance") or {}).get("checked")],
            ["Running balance breaks", (bc.get("running_balance") or {}).get("breaks_count"), (ledger.get("running_balance") or {}).get("breaks_count")],
            [],
            ["Book-to-bank check", btb.get("status")],
            ["Closing balance difference (bank - ledger)", btb.get("closing_balance_difference")],
            ["Opening balance difference (bank - ledger)", btb.get("opening_balance_difference")],
            ["Explained by discrepancies", btb.get("explained_by_discrepancies")],
            ["Unexplained difference", btb.get("unexplained_difference")],
            [],
            ["Running balance breaks", "Transaction", "Date", "Expected", "Stated", "Difference"],
        ]
        for source, rb in (("Bank", bc.get("running_balance") or {}), ("Ledger", ledger.get("running_balance") or {})):
            for brk in rb.get("breaks", []):
                rows.append([source, brk["transaction_id"], brk["date"], brk["expected_balance"], brk["stated_balance"], brk["difference"]])
        return rows

    @staticmethod
    def _section_tables(result: Dict[str, Any]) -> List[Tuple[str, List[str], List[List[Any]]]]:
        def tx(t, *keys):
            return [t.get(k) for k in keys]

        matched = [
            tx(m["bank_transaction"], "id", "date", "description", "reference")
            + tx(m["ledger_transaction"], "id", "date", "description", "reference")
            + [m["amount"], m["bank_transaction"]["flow"], m.get("date_difference_days"), ", ".join(m["match_basis"]), m["confidence"]]
            for m in result["matched_transactions"]
        ]
        unmatched_cols = ["id", "row", "date", "description", "reference", "debit", "credit", "balance", "flow", "reason"]
        unmatched_bank = [tx(t, *unmatched_cols) for t in result["unmatched_bank_transactions"]]
        unmatched_ledger = [tx(t, *unmatched_cols) for t in result["unmatched_ledger_transactions"]]
        duplicates = [
            [d["source"], ", ".join(d["transaction_ids"]), d["transactions"][0]["date"], d["transactions"][0]["amount"], d["count"], d["reason"]]
            for d in result["possible_duplicates"]
        ]
        mismatches = [
            [m["type"]]
            + tx(m["bank_transaction"], "id", "date", "description")
            + [m["bank_amount"]]
            + tx(m["ledger_transaction"], "id", "date", "description")
            + [m["ledger_amount"], m["difference"], m["reason"]]
            for m in result["amount_mismatches"]
        ]
        return [
            ("Matched", ["Bank ID", "Bank Date", "Bank Description", "Bank Reference", "Ledger ID", "Ledger Date",
                         "Ledger Description", "Ledger Reference", "Amount", "Flow", "Date Diff (days)", "Match Basis", "Confidence"], matched),
            ("Unmatched Bank", ["ID", "Row", "Date", "Description", "Reference", "Debit", "Credit", "Balance", "Flow", "Reason"], unmatched_bank),
            ("Unmatched Ledger", ["ID", "Row", "Date", "Description", "Reference", "Debit", "Credit", "Balance", "Flow", "Reason"], unmatched_ledger),
            ("Possible Duplicates", ["Source", "Transaction IDs", "Date", "Amount", "Count", "Reason"], duplicates),
            ("Amount Mismatches", ["Type", "Bank ID", "Bank Date", "Bank Description", "Bank Amount", "Ledger ID", "Ledger Date",
                                   "Ledger Description", "Ledger Amount", "Difference", "Reason"], mismatches),
            ("Amount Corrections", ["Source", "Transaction ID", "Field", "Extracted", "Corrected"], [
                [c["source"], c["transaction_id"], c["field"], c["extracted"], c["corrected"]]
                for c in result.get("amount_corrections", [])
            ]),
        ]

    @classmethod
    def build_export(cls, record: Dict[str, Any], fmt: str) -> Tuple[bytes, str, str]:
        """Builds an XLSX or CSV report from the saved result. Returns (content, media_type, filename)."""
        result = record["result"]
        sections = cls._section_tables(result)
        warnings = [[w] for w in result.get("warnings", [])]
        rid = record["reconciliation_id"]

        if fmt == "csv":
            buffer = io.StringIO()
            writer = csv.writer(buffer)
            blocks = [("Summary", None, cls._summary_rows(record)), ("Balance Verification", None, cls._balance_rows(result))]
            blocks += sections
            blocks.append(("Warnings", ["Warning"], warnings))
            for title, header, rows in blocks:
                writer.writerow([f"## {title}"])
                if header:
                    writer.writerow(header)
                for row in rows:
                    writer.writerow([sanitize_for_formula_injection(c) for c in row])
                writer.writerow([])
            return buffer.getvalue().encode("utf-8-sig"), "text/csv", f"{rid}.csv"

        import openpyxl
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F46E5")
        amount_headers = {"Amount", "Debit", "Credit", "Balance", "Bank Amount", "Ledger Amount", "Difference", "Extracted", "Corrected"}
        places = int((result.get("summary") or {}).get("amount_decimal_places", 2))
        number_format = "#,##0" + ("." + "0" * places if places else "")

        def cell_value(value: Any, numeric: bool) -> Any:
            if numeric and isinstance(value, str):
                try:
                    return float(Decimal(value))
                except Exception:
                    pass
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)  # OCR text can contain control characters
            return sanitize_for_formula_injection(value)

        def write_sheet(ws, header: Optional[List[str]], rows: List[List[Any]]) -> None:
            numeric_cols = set()
            if header:
                ws.append(header)
                for cell in ws[1]:
                    cell.font, cell.fill = header_font, header_fill
                numeric_cols = {i for i, h in enumerate(header) if h in amount_headers}
                ws.freeze_panes = "A2"
            for row in rows:
                ws.append([cell_value(v, i in numeric_cols) for i, v in enumerate(row)])
                if numeric_cols:
                    for i in numeric_cols:
                        if i < len(row):
                            ws.cell(row=ws.max_row, column=i + 1).number_format = number_format
            for idx in range(1, (len(header) if header else 6) + 1):
                ws.column_dimensions[get_column_letter(idx)].width = 22

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        write_sheet(ws, None, cls._summary_rows(record))
        for cell in ws["A"]:
            cell.font = Font(bold=True)
        write_sheet(wb.create_sheet("Balance Verification"), None, cls._balance_rows(result))
        for title, header, rows in sections:
            write_sheet(wb.create_sheet(title), header, rows)
        write_sheet(wb.create_sheet("Warnings"), ["Warning"], warnings)

        out = io.BytesIO()
        wb.save(out)
        return (
            out.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{rid}.xlsx",
        )
