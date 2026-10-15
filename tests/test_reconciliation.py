"""
Tests for the Bank Reconciliation engine and its four API endpoints.
"""

import asyncio
import hashlib
import hmac
import io
import json
import socket
from datetime import datetime, timedelta, timezone

import openpyxl
import pymupdf
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient

from app.core.config import get_settings
from app.db.mongodb import MongoDBManager
from app.main import app
from app.schemas.ocr import ExtractionResponse, ProcessingMetadata
from app.services import reconciliation_service as rec_service_module
from app.services.pipeline import ExtractionPipeline
from app.services.reconciliation_engine import (
    ReconciliationError,
    ReconciliationOptions,
    parse_spreadsheet,
    reconcile,
)
from app.services.reconciliation_service import ReconciliationService
from app.services.webhook_service import WebhookService
from app.workers.worker_manager import WorkerManager

BASE = "/api/v1/reconciliation"


# ---------------------------------------------------------------------------
# Fixtures & helpers
# ---------------------------------------------------------------------------

@pytest.fixture
def api_key():
    return get_settings().api_keys[0]


@pytest_asyncio.fixture(autouse=True)
async def environment(tmp_path, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "storage_dir", str(tmp_path))
    monkeypatch.setattr(settings, "rate_limit_enabled", False)
    monkeypatch.setattr(settings, "require_api_key", True)
    await MongoDBManager.disconnect()
    await MongoDBManager.connect(force_mock=True)
    await WorkerManager.start()
    yield
    await WorkerManager.stop()
    await MongoDBManager.disconnect()


@pytest.fixture
def public_dns(monkeypatch):
    """Resolve every webhook host to a public IP so SSRF checks pass without network access."""
    def fake_getaddrinfo(host, port, *args, **kwargs):
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", port))]
    monkeypatch.setattr(rec_service_module.socket, "getaddrinfo", fake_getaddrinfo)


@pytest.fixture
def captured_webhooks(monkeypatch):
    sent = []

    async def fake_send(cls, callback_url, payload, secret=None, max_retries=3, event_id=None, job_id=None, event_name=None):
        sent.append({"url": callback_url, "payload": payload, "secret": secret, "event": event_name})
        return True

    monkeypatch.setattr(WebhookService, "send_webhook", classmethod(fake_send))
    return sent


def client():
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


BANK_CSV = (
    "Date,Narration,Chq/Ref No,Withdrawal Amt,Deposit Amt,Balance\n"
    "01/10/2026,Opening Balance,,,,10000.00\n"
    "01/10/2026,NEFT ACME Corp invoice 101,UTR1001,,5000.00,15000.00\n"
    "03/10/2026,Office rent October,CHQ2001,3000.00,,12000.00\n"
    "05/10/2026,Bank charges,,50.00,,11950.00\n"
    "07/10/2026,Payment to Supplier Beta,UTR1003,1200.00,,10750.00\n"
)
LEDGER_CSV = (
    "Date,Particulars,Voucher No,Debit,Credit,Balance\n"
    "01/10/2026,Opening Balance,,,,10000.00\n"
    "02/10/2026,ACME Corp invoice 101 receipt,UTR1001,,5000.00,15000.00\n"
    "03/10/2026,Rent for October office,CHQ2001,3000.00,,12000.00\n"
    "07/10/2026,Supplier Beta payment,UTR1003,1250.00,,10750.00\n"
    "09/10/2026,Cash deposit,,,700.00,11450.00\n"
)


def files(bank=("bank.csv", BANK_CSV.encode(), "text/csv"), ledger=("ledger.csv", LEDGER_CSV.encode(), "text/csv")):
    return {"bank_statement": bank, "business_ledger": ledger}


async def wait_for(c, rid, headers, terminal=("completed", "failed"), attempts=200):
    body = None
    for _ in range(attempts):
        resp = await c.get(f"{BASE}/{rid}", headers=headers)
        assert resp.status_code == 200, resp.text
        body = resp.json()
        if body["status"] in terminal:
            return body
        await asyncio.sleep(0.02)
    raise AssertionError(f"Reconciliation never reached {terminal}: {body}")


def make_pdf(pages: int = 1) -> bytes:
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page().insert_text((50, 50), f"Statement page {i + 1}")
    data = doc.write()
    doc.close()
    return data


def tx(date, desc, debit=None, credit=None, ref=None, balance=None):
    return {"date": date, "description": desc, "reference": ref, "debit": debit, "credit": credit, "balance": balance}


# ---------------------------------------------------------------------------
# Engine
# ---------------------------------------------------------------------------

def test_engine_matches_with_tolerance_and_reports_discrepancies():
    bank = {
        "currency": "INR", "opening_balance": 10000.0, "closing_balance": 10750.0,
        "transactions": [
            tx("2026-10-01", "NEFT ACME Corp", credit=5000, ref="UTR1", balance=15000),
            tx("2026-10-03", "Office rent", debit=3000, balance=12000),
            tx("2026-10-05", "Bank charges", debit=50, balance=11950),
            tx("2026-10-07", "Supplier Beta", debit=1200, ref="UTR3", balance=10750),
        ],
    }
    ledger = {
        "currency": "INR", "opening_balance": 10000.0, "closing_balance": 11450.0,
        "transactions": [
            tx("2026-10-02", "ACME Corp receipt", credit=5000, ref="UTR1"),
            tx("2026-10-03", "Rent office", debit=3000),
            tx("2026-10-07", "Supplier Beta", debit=1250, ref="UTR3"),
            tx("2026-10-09", "Cash deposit", credit=700),
        ],
    }
    result = reconcile(bank, ledger, ReconciliationOptions(date_tolerance_days=2))
    s = result["summary"]
    assert (s["matched"], s["unmatched_bank"], s["unmatched_ledger"], s["amount_mismatches"]) == (2, 1, 1, 1)
    assert s["currency"] == "INR"
    assert result["unmatched_bank_transactions"][0]["description"] == "Bank charges"
    assert result["unmatched_ledger_transactions"][0]["description"] == "Cash deposit"
    mismatch = result["amount_mismatches"][0]
    assert mismatch["type"] == "amount_mismatch" and mismatch["difference"] == "50.00"
    acme = next(m for m in result["matched_transactions"] if m["bank_transaction"]["id"] == "B1")
    assert acme["date_difference_days"] == 1 and "reference" in acme["match_basis"]

    bc = result["balance_check"]
    assert bc["status"] == "passed" and bc["difference"] == "0.00"
    assert bc["calculated_closing_balance"] == "10750.00"
    assert bc["running_balance"]["status"] == "passed"
    # Bank 10750 - ledger 11450 = -700 = -50 (bank charges) - 700 (cash deposit) + 50 (mismatch)
    assert bc["book_to_bank"]["status"] == "passed"
    assert bc["book_to_bank"]["unexplained_difference"] == "0.00"
    assert any("not, by themselves, evidence of fraud" in w for w in result["warnings"])


def test_engine_date_tolerance_is_configurable():
    bank = {"transactions": [tx("2026-10-01", "Payment X", debit=100)]}
    ledger = {"transactions": [tx("2026-10-05", "Payment X", debit=100)]}
    assert reconcile(bank, ledger, ReconciliationOptions(date_tolerance_days=2))["summary"]["matched"] == 0
    narrow = reconcile(bank, ledger, ReconciliationOptions(date_tolerance_days=2))
    assert "outside the ±2-day tolerance" in narrow["unmatched_bank_transactions"][0]["reason"]
    assert reconcile(bank, ledger, ReconciliationOptions(date_tolerance_days=4))["summary"]["matched"] == 1


def test_engine_never_matches_a_transaction_twice_and_flags_duplicates():
    bank = {"transactions": [
        tx("2026-10-01", "Coffee supplies", debit=99.99),
        tx("2026-10-01", "Coffee supplies", debit=99.99),
    ]}
    ledger = {"transactions": [tx("2026-10-01", "Coffee supplies", debit=99.99)]}
    result = reconcile(bank, ledger)
    assert result["summary"]["matched"] == 1
    assert result["summary"]["unmatched_bank"] == 1
    assert result["summary"]["possible_duplicates"] == 1
    dup = result["possible_duplicates"][0]
    assert dup["source"] == "bank" and dup["transaction_ids"] == ["B1", "B2"]
    # Original transactions are preserved, nothing deleted
    assert result["summary"]["bank_transactions"] == 2


def test_engine_detects_direction_mismatch():
    bank = {"transactions": [tx("2026-10-02", "Refund from vendor", credit=400, ref="R77"),
                             tx("2026-10-03", "Salary", debit=1000)]}
    ledger = {"transactions": [tx("2026-10-02", "Refund vendor", debit=400, ref="R77"),
                               tx("2026-10-03", "Salary", debit=1000)]}
    result = reconcile(bank, ledger)
    assert result["summary"]["matched"] == 1
    assert result["amount_mismatches"][0]["type"] == "direction_mismatch"
    assert result["summary"]["ledger_direction_convention"] == "bank"


def test_engine_detects_book_keeping_convention():
    # In the company's books a bank deposit is a debit; the engine should detect and adapt.
    bank = {"transactions": [tx("2026-10-01", "Deposit", credit=500), tx("2026-10-02", "Rent", debit=200)]}
    ledger = {"transactions": [tx("2026-10-01", "Deposit", debit=500), tx("2026-10-02", "Rent", credit=200)]}
    result = reconcile(bank, ledger)
    assert result["summary"]["matched"] == 2
    assert result["summary"]["ledger_direction_convention"] == "book"
    assert any("book-keeping convention" in w for w in result["warnings"])


def test_engine_balance_failures_and_unverifiable_data():
    bank = {"opening_balance": 1000, "closing_balance": 1500, "transactions": [
        tx("2026-10-01", "In", credit=300, balance=1300),
        tx("2026-10-02", "In", credit=100, balance=1450),  # should be 1400
    ]}
    ledger = {"transactions": [tx("2026-10-01", "In", credit=300), tx("2026-10-02", "In", credit=100)]}
    bc = reconcile(bank, ledger)["balance_check"]
    assert bc["status"] == "failed" and bc["difference"] == "100.00"
    assert bc["running_balance"]["breaks_count"] == 1
    assert bc["running_balance"]["breaks"][0]["expected_balance"] == "1400.00"
    assert bc["ledger"]["status"] == "unverifiable"
    assert bc["book_to_bank"]["status"] == "unverifiable"


def test_engine_uses_decimal_precision():
    bank = {"opening_balance": "0.10", "closing_balance": "0.30", "transactions": [tx("2026-10-01", "a", credit=0.2)]}
    ledger = {"transactions": [tx("2026-10-01", "a", credit="0.20")]}
    result = reconcile(bank, ledger)
    assert result["balance_check"]["difference"] == "0.00"
    assert result["summary"]["matched"] == 1


def test_engine_rejects_currency_mismatch_and_empty_documents():
    with pytest.raises(ReconciliationError, match="cross-currency"):
        reconcile({"currency": "INR", "transactions": [tx("2026-10-01", "a", credit=1)]},
                  {"currency": "USD", "transactions": [tx("2026-10-01", "a", credit=1)]})
    with pytest.raises(ReconciliationError, match="business ledger"):
        reconcile({"transactions": [tx("2026-10-01", "a", credit=1)]}, {"transactions": []})


def test_spreadsheet_parser_csv_and_xlsx():
    parsed = parse_spreadsheet(BANK_CSV.encode(), "bank.csv")
    assert parsed["opening_balance"] == "10000.00"
    assert len(parsed["transactions"]) == 4

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Company Ledger"])
    ws.append(["Txn Date", "Description", "Amount", "Dr/Cr", "Balance"])
    ws.append([datetime(2026, 10, 1), "Opening Balance", None, None, 500])
    ws.append([datetime(2026, 10, 2), "Sale", "1,200.50", "Cr", 1700.5])
    ws.append([datetime(2026, 10, 3), "Purchase", 200, "Dr", 1500.5])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_spreadsheet(buf.getvalue(), "ledger.xlsx")
    assert parsed["opening_balance"] == 500
    result = reconcile(
        {"transactions": [tx("2026-10-02", "Sale", credit=1200.5), tx("2026-10-03", "Purchase", debit=200)]},
        parsed,
    )
    assert result["summary"]["matched"] == 2

    with pytest.raises(ReconciliationError, match="header row"):
        parse_spreadsheet(b"foo,bar\n1,2\n", "x.csv")


# ---------------------------------------------------------------------------
# API 1 + 2: upload & status (guest and authenticated)
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_guest_upload_and_retrieval(api_key):
    async with client() as c:
        resp = await c.post(f"{BASE}/upload", files=files(), data={"date_tolerance_days": "2"})
        assert resp.status_code == 202, resp.text
        body = resp.json()
        rid, token = body["reconciliation_id"], body["access_token"]
        assert body["status"] == "queued" and body["status_url"] == f"/api/v1/reconciliation/{rid}"
        assert body["success"] is True and len(token) >= 40

        final = await wait_for(c, rid, {"X-Reconciliation-Token": token})
        assert final["status"] == "completed" and final["progress"] == 100
        assert final["result"]["summary"]["matched"] == 2
        assert final["result"]["summary"]["unmatched_bank"] == 1
        assert final["result"]["summary"]["amount_mismatches"] == 1
        assert final["result"]["balance_check"]["status"] == "passed"
        assert final["completed_at"] and final["error"] is None

        # No token, wrong token, or a valid API key all fail to read a guest report
        assert (await c.get(f"{BASE}/{rid}")).status_code == 404
        assert (await c.get(f"{BASE}/{rid}", headers={"X-Reconciliation-Token": "guess"})).status_code == 404
        assert (await c.get(f"{BASE}/{rid}", headers={"X-API-Key": api_key})).status_code == 404

    # Only the token hash is persisted
    record = await MongoDBManager.get_db()["reconciliations"].find_one({"reconciliation_id": rid})
    assert token not in json.dumps(record, default=str)
    assert record["access_token_hash"] == hashlib.sha256(token.encode()).hexdigest()
    assert record["owner_id"] is None and record["expires_at"] is not None


@pytest.mark.asyncio
async def test_authenticated_upload_is_owner_scoped(api_key, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "api_keys", [api_key, "second_tenant_key_123"])
    async with client() as c:
        resp = await c.post(f"{BASE}/upload", files=files(), data={"user_email": "Owner@Example.com"},
                            headers={"Authorization": f"Bearer {api_key}"})
        assert resp.status_code == 202
        body = resp.json()
        assert "access_token" not in body
        rid = body["reconciliation_id"]

        final = await wait_for(c, rid, {"X-API-Key": api_key})
        assert final["status"] == "completed"
        assert (await c.get(f"{BASE}/{rid}", headers={"X-API-Key": "second_tenant_key_123"})).status_code == 404
        assert (await c.get(f"{BASE}/{rid}")).status_code == 404
        assert (await c.get(f"{BASE}/{rid}", headers={"X-API-Key": "invalid-key"})).status_code == 401

    record = await MongoDBManager.get_db()["reconciliations"].find_one({"reconciliation_id": rid})
    assert record["owner_email"] == "owner@example.com" and record["access_token_hash"] is None


@pytest.mark.asyncio
async def test_invalid_ids_and_expired_guest_access():
    async with client() as c:
        assert (await c.get(f"{BASE}/../etc")).status_code == 404
        assert (await c.get(f"{BASE}/rec_deadbeefdeadbeef")).status_code == 404
        assert (await c.get(f"{BASE}/not-an-id/transactions")).status_code == 404

        body = (await c.post(f"{BASE}/upload", files=files())).json()
        rid, headers = body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]}
        await wait_for(c, rid, headers)
        await MongoDBManager.get_db()["reconciliations"].update_one(
            {"reconciliation_id": rid}, {"$set": {"expires_at": datetime.now(timezone.utc) - timedelta(minutes=1)}}
        )
        assert (await c.get(f"{BASE}/{rid}", headers=headers)).status_code == 410
        assert (await c.get(f"{BASE}/{rid}/export", headers=headers)).status_code == 410


@pytest.mark.asyncio
async def test_upload_validation(public_dns):
    async with client() as c:
        bad_ext = await c.post(f"{BASE}/upload", files=files(ledger=("ledger.docx", b"abc", "application/msword")))
        assert bad_ext.status_code == 400 and "business ledger" in bad_ext.json()["error"]

        empty = await c.post(f"{BASE}/upload", files=files(bank=("bank.csv", b"", "text/csv")))
        assert empty.status_code == 400

        fake_xlsx = await c.post(f"{BASE}/upload", files=files(ledger=("ledger.xlsx", b"not a zip", "application/octet-stream")))
        assert fake_xlsx.status_code == 400

        fake_pdf = await c.post(f"{BASE}/upload", files=files(bank=("bank.pdf", b"hello", "application/pdf")))
        assert fake_pdf.status_code == 400

        missing = await c.post(f"{BASE}/upload", files={"bank_statement": ("bank.csv", BANK_CSV.encode(), "text/csv")})
        assert missing.status_code == 422

        bad_tolerance = await c.post(f"{BASE}/upload", files=files(), data={"date_tolerance_days": "99"})
        assert bad_tolerance.status_code == 422

        http_hook = await c.post(f"{BASE}/upload", files=files(), data={"webhook_url": "http://example.com/hook"})
        assert http_hook.status_code == 400 and "HTTPS" in http_hook.json()["error"]

        local_hook = await c.post(f"{BASE}/upload", files=files(), data={"webhook_url": "https://localhost/hook"})
        assert local_hook.status_code == 400


@pytest.mark.parametrize("ip", ["127.0.0.1", "10.0.0.5", "169.254.169.254", "192.168.1.1", "::1"])
def test_webhook_url_ssrf_protection(monkeypatch, ip):
    family = socket.AF_INET6 if ":" in ip else socket.AF_INET
    monkeypatch.setattr(rec_service_module.socket, "getaddrinfo",
                        lambda *a, **k: [(family, socket.SOCK_STREAM, 6, "", (ip, 443))])
    with pytest.raises(ValueError, match="private or internal"):
        ReconciliationService.validate_webhook_url("https://hooks.example.com/x")


# ---------------------------------------------------------------------------
# Large documents through the existing OCR worker pipeline
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_large_pdf_processed_asynchronously_via_ocr_jobs(monkeypatch, api_key, public_dns, captured_webhooks):
    release = asyncio.Event()

    async def fake_process_document(self, **kwargs):
        await release.wait()
        extraction = {
            "currency": "INR", "opening_balance": 10000.0, "closing_balance": 10750.0,
            "transactions": [
                tx("2026-10-01", "NEFT ACME Corp invoice 101", credit=5000.0, ref="UTR1001", balance=15000.0),
                tx("2026-10-03", "Office rent October", debit=3000.0, ref="CHQ2001", balance=12000.0),
                tx("2026-10-05", "Bank charges", debit=50.0, balance=11950.0),
                tx("2026-10-07", "Payment to Supplier Beta", debit=1200.0, ref="UTR1003", balance=10750.0),
            ],
        }
        return ExtractionResponse(id=kwargs["client_request_id"], document_type="bank_statement", extraction=extraction,
                                  raw_text="...", metadata=ProcessingMetadata(pages=150))

    monkeypatch.setattr(ExtractionPipeline, "process_document", fake_process_document)

    async with client() as c:
        resp = await c.post(
            f"{BASE}/upload",
            files=files(bank=("statement.pdf", make_pdf(150), "application/pdf")),
            data={"webhook_url": "https://hooks.example.com/rec", "webhook_secret": "s3cret"},
            headers={"X-API-Key": api_key},
        )
        assert resp.status_code == 202
        rid = resp.json()["reconciliation_id"]

        record = await MongoDBManager.get_db()["reconciliations"].find_one({"reconciliation_id": rid})
        job_id = record["documents"]["bank_statement"]["job_id"]
        job = await MongoDBManager.get_db()["jobs"].find_one({"job_id": job_id})
        assert job["metadata"]["reconciliation_id"] == rid and job["document_type"] == "bank_statement"
        assert record["documents"]["business_ledger"]["job_id"] is None

        processing = None
        for _ in range(100):
            processing = (await c.get(f"{BASE}/{rid}", headers={"X-API-Key": api_key})).json()
            if processing["status"] == "processing":
                break
            await asyncio.sleep(0.01)
        assert processing["status"] == "processing"
        assert processing["message"] == "Extracting bank statement transactions"
        assert 0 < processing["progress"] < 100 and processing["result"] is None

        release.set()
        final = await wait_for(c, rid, {"X-API-Key": api_key})
        assert final["status"] == "completed"
        assert final["result"]["summary"]["matched"] == 2

    for _ in range(100):
        if captured_webhooks:
            break
        await asyncio.sleep(0.01)
    assert captured_webhooks[0]["event"] == "reconciliation.completed"
    payload = captured_webhooks[0]["payload"]
    assert payload["reconciliation_id"] == rid and payload["result_url"] == f"/api/v1/reconciliation/{rid}"
    assert captured_webhooks[0]["secret"] == "s3cret"

    await asyncio.sleep(0.05)
    record = await MongoDBManager.get_db()["reconciliations"].find_one({"reconciliation_id": rid})
    assert record["webhook"]["delivery_status"] == "delivered"
    assert "secret" not in record["webhook"]


@pytest.mark.asyncio
async def test_ocr_failure_marks_reconciliation_failed(monkeypatch, public_dns, captured_webhooks):
    async def failing_process_document(self, **kwargs):
        raise ValueError("Document is unreadable")

    monkeypatch.setattr(ExtractionPipeline, "process_document", failing_process_document)
    async with client() as c:
        body = (await c.post(
            f"{BASE}/upload",
            files=files(ledger=("ledger.pdf", make_pdf(2), "application/pdf")),
            data={"webhook_url": "https://hooks.example.com/rec"},
        )).json()
        rid, headers = body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]}
        final = await wait_for(c, rid, headers)
        assert final["status"] == "failed" and final["success"] is False
        assert final["error"] == "Unable to extract transactions from the business ledger."
        assert final["result"] is None

        assert (await c.get(f"{BASE}/{rid}/transactions", headers=headers)).status_code == 409

    for _ in range(100):
        if captured_webhooks:
            break
        await asyncio.sleep(0.01)
    assert captured_webhooks[0]["payload"]["event"] == "reconciliation.failed"
    assert captured_webhooks[0]["payload"]["error"] == "Unable to extract transactions from the business ledger."


@pytest.mark.asyncio
async def test_webhook_signature_and_bounded_retries(monkeypatch):
    calls = []

    class FakeResponse:
        def __init__(self, code):
            self.status_code, self.text = code, ""
            self.is_success = 200 <= code < 300

    async def fake_post(self, url, content=None, headers=None):
        calls.append((content, headers))
        return FakeResponse(500 if len(calls) < 3 else 200)

    async def no_sleep(_):
        return None

    import httpx
    monkeypatch.setattr(httpx.AsyncClient, "post", fake_post)
    monkeypatch.setattr("app.services.webhook_service.asyncio.sleep", no_sleep)

    payload = {"event": "reconciliation.completed", "reconciliation_id": "rec_abc12345", "status": "completed"}
    ok = await WebhookService.send_webhook("https://hooks.example.com/x", payload, secret="topsecret",
                                           max_retries=rec_service_module.WEBHOOK_MAX_ATTEMPTS, job_id="rec_abc12345",
                                           event_name="reconciliation.completed")
    assert ok and len(calls) == 3
    body, headers = calls[-1]
    expected = hmac.new(b"topsecret", body, hashlib.sha256).hexdigest()
    assert headers["X-Webhook-Signature"] == f"sha256={expected}"


@pytest.mark.asyncio
async def test_finalize_is_idempotent_and_get_does_not_reprocess(monkeypatch):
    calls = {"n": 0}
    real_reconcile = rec_service_module.reconcile

    def counting_reconcile(*args, **kwargs):
        calls["n"] += 1
        return real_reconcile(*args, **kwargs)

    monkeypatch.setattr(rec_service_module, "reconcile", counting_reconcile)
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files())).json()
        rid, headers = body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]}
        await wait_for(c, rid, headers)
        await asyncio.gather(*(ReconciliationService.try_finalize(rid) for _ in range(5)))
        for _ in range(3):
            assert (await c.get(f"{BASE}/{rid}", headers=headers)).json()["status"] == "completed"
            await c.get(f"{BASE}/{rid}/export?format=csv", headers=headers)
    assert calls["n"] == 1
    assert await MongoDBManager.get_db()["reconciliations"].count_documents({"reconciliation_id": rid}) == 1


# ---------------------------------------------------------------------------
# API 3: transactions
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_transactions_filter_and_pagination():
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files())).json()
        rid, headers = body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]}
        await wait_for(c, rid, headers)

        matched = (await c.get(f"{BASE}/{rid}/transactions?status=matched", headers=headers)).json()
        assert matched["total"] == 2 and len(matched["transactions"]) == 2
        assert all(t["category"] == "matched" for t in matched["transactions"])

        page1 = (await c.get(f"{BASE}/{rid}/transactions?page=1&page_size=2", headers=headers)).json()
        page2 = (await c.get(f"{BASE}/{rid}/transactions?page=2&page_size=2", headers=headers)).json()
        page3 = (await c.get(f"{BASE}/{rid}/transactions?page=3&page_size=2", headers=headers)).json()
        assert page1["total"] == 5  # 2 matched + 1 unmatched bank + 1 unmatched ledger + 1 mismatch
        assert [t["category"] for t in page1["transactions"]] == ["matched", "matched"]
        assert [t["category"] for t in page2["transactions"]] == ["unmatched_bank", "unmatched_ledger"]
        assert [t["category"] for t in page3["transactions"]] == ["amount_mismatches"]

        mismatch = (await c.get(f"{BASE}/{rid}/transactions?status=amount_mismatches", headers=headers)).json()
        assert mismatch["transactions"][0]["difference"] == "50.00"

        assert (await c.get(f"{BASE}/{rid}/transactions?status=bogus", headers=headers)).status_code == 422
        assert (await c.get(f"{BASE}/{rid}/transactions?page_size=0", headers=headers)).status_code == 422
        assert (await c.get(f"{BASE}/{rid}/transactions?status=matched")).status_code == 404


# ---------------------------------------------------------------------------
# API 4: export
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_export_xlsx_and_csv_with_formula_sanitization():
    malicious_ledger = LEDGER_CSV.replace("Cash deposit", "=HYPERLINK(\"http://evil\")")
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files(ledger=("ledger.csv", malicious_ledger.encode(), "text/csv")))).json()
        rid, headers = body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]}
        await wait_for(c, rid, headers)

        resp = await c.get(f"{BASE}/{rid}/export?format=xlsx", headers=headers)
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert f'filename="{rid}.xlsx"' in resp.headers["content-disposition"]
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["Summary", "Balance Verification", "Matched", "Unmatched Bank", "Unmatched Ledger",
                                 "Possible Duplicates", "Amount Mismatches", "Amount Corrections", "Warnings"]
        ledger_rows = list(wb["Unmatched Ledger"].iter_rows(values_only=True))
        assert ledger_rows[1][3].startswith("'=")
        assert wb["Matched"].cell(row=2, column=9).value == 5000.0

        csv_resp = await c.get(f"{BASE}/{rid}/export?format=csv", headers=headers)
        assert csv_resp.status_code == 200 and csv_resp.headers["content-type"].startswith("text/csv")
        text = csv_resp.content.decode("utf-8-sig")
        for section in ("## Summary", "## Balance Verification", "## Matched", "## Unmatched Bank",
                        "## Unmatched Ledger", "## Possible Duplicates", "## Amount Mismatches",
                        "## Amount Corrections", "## Warnings"):
            assert section in text
        assert "'=HYPERLINK" in text

        assert (await c.get(f"{BASE}/{rid}/export?format=pdf", headers=headers)).status_code == 422
        assert (await c.get(f"{BASE}/{rid}/export")).status_code == 404


# ---------------------------------------------------------------------------
# Audit regressions
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("raw,expected,marker", [
    ("Rs. 500", "500.00", None),
    ("INR 1,23,456.78", "123456.78", None),
    ("500.00Cr", "500.00", "cr"),
    ("12,000.00 Dr", "12000.00", "dr"),
    ("1,234.50-", "-1234.50", None),
    ("(250.00)", "-250.00", None),
    ("−75.10", "-75.10", None),
    ("12-34", None, None),
    ("", None, None),
])
def test_parse_amount_formats(raw, expected, marker):
    from app.services.reconciliation_engine import parse_amount
    amount, got_marker = parse_amount(raw)
    assert (str(amount) if amount is not None else None) == expected
    assert got_marker == marker


def test_month_first_dates_detected_per_document():
    bank = {"transactions": [tx("2026-10-03", "Rent", debit=100), tx("2026-10-13", "Fees", debit=20)]}
    ledger = {"transactions": [tx("10/03/2026", "Rent", debit=100), tx("10/13/2026", "Fees", debit=20)]}
    result = reconcile(bank, ledger)
    assert result["summary"]["matched"] == 2
    assert any("month-first" in w for w in result["warnings"])


def test_unrelated_same_day_opposite_entries_are_not_direction_mismatches():
    bank = {"transactions": [tx("2026-10-01", "Customer receipt Alpha", credit=500)]}
    ledger = {"transactions": [tx("2026-10-01", "Stationery purchase", debit=500)]}
    result = reconcile(bank, ledger)
    assert result["summary"]["amount_mismatches"] == 0
    assert result["summary"]["unmatched_bank"] == 1 and result["summary"]["unmatched_ledger"] == 1


def test_overdraft_dr_balances_verify_correctly():
    bank = {"opening_balance": "100.00", "transactions": [
        tx("2026-10-01", "Payment", debit=300, balance="200.00 Dr"),
        tx("2026-10-02", "Deposit", credit=500, balance="300.00 Cr"),
    ]}
    ledger = {"transactions": [tx("2026-10-01", "Payment", debit=300), tx("2026-10-02", "Deposit", credit=500)]}
    bc = reconcile(bank, ledger)["balance_check"]
    assert bc["running_balance"]["status"] == "passed"
    assert bc["status"] == "passed" and bc["bank_closing_balance"] == "300.00"


def test_spreadsheet_header_priority_and_opening_row_detection():
    csv_text = (
        "Value Date,Txn Date,Description,Debit,Credit\n"
        "05/10/2026,01/10/2026,Opening Balance,,1000\n"
        "06/10/2026,02/10/2026,Transfer to opening balance adjustment,50,\n"
    )
    parsed = parse_spreadsheet(csv_text.encode(), "s.csv")
    assert parsed["opening_balance"] == "1000"
    assert len(parsed["transactions"]) == 1
    assert parsed["transactions"][0]["date"] == "02/10/2026"


def test_large_identical_amount_volume_is_fast():
    import time
    bank = {"transactions": [tx(f"2026-10-{(i % 28) + 1:02d}", "Fee", debit=100) for i in range(3000)]}
    ledger = {"transactions": [tx(f"2026-10-{(i % 28) + 1:02d}", "Fee", debit=100) for i in range(3000)]}
    started = time.perf_counter()
    result = reconcile(bank, ledger, ReconciliationOptions(description_matching=False))
    assert result["summary"]["matched"] == 3000
    assert time.perf_counter() - started < 20


@pytest.mark.asyncio
async def test_pdf_over_page_limit_rejected_at_upload(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_pdf_pages", 2)
    async with client() as c:
        resp = await c.post(f"{BASE}/upload", files=files(bank=("bank.pdf", make_pdf(3), "application/pdf")))
        assert 400 <= resp.status_code < 500
        assert resp.json()["error"].startswith("Bank statement:")
    assert await MongoDBManager.get_db()["jobs"].count_documents({}) == 0


@pytest.mark.asyncio
async def test_failed_record_insert_cleans_up_jobs(monkeypatch):
    from app.db.repositories.reconciliation_repo import ReconciliationRepository

    async def boom(cls, record):
        raise RuntimeError("db down")

    monkeypatch.setattr(ReconciliationRepository, "create", classmethod(boom))
    with pytest.raises(RuntimeError):
        await ReconciliationService.create_reconciliation(
            documents={
                "bank_statement": ("pdf", "bank.pdf", make_pdf(1), "application/pdf"),
                "business_ledger": ("spreadsheet", "ledger.csv", LEDGER_CSV.encode(), "text/csv"),
            },
            options=ReconciliationOptions(), owner_id=None, owner_email=None, webhook_url=None, webhook_secret=None,
        )
    assert await MongoDBManager.get_db()["jobs"].count_documents({}) == 0


@pytest.mark.asyncio
async def test_oversized_report_fails_cleanly(monkeypatch):
    from pymongo.errors import DocumentTooLarge
    from app.db.repositories.reconciliation_repo import ReconciliationRepository

    async def too_large(cls, reconciliation_id, result):
        raise DocumentTooLarge("too big")

    monkeypatch.setattr(ReconciliationRepository, "complete", classmethod(too_large))
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files())).json()
        final = await wait_for(c, body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]})
    assert final["status"] == "failed" and "too large" in final["error"]


@pytest.mark.asyncio
async def test_export_strips_control_characters():
    dirty = LEDGER_CSV.replace("Cash deposit", "Cash\x0b dep\x01osit")
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files(ledger=("ledger.csv", dirty.encode(), "text/csv")))).json()
        rid, headers = body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]}
        await wait_for(c, rid, headers)
        resp = await c.get(f"{BASE}/{rid}/export?format=xlsx", headers=headers)
        assert resp.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert list(wb["Unmatched Ledger"].iter_rows(values_only=True))[1][3] == "Cash deposit"


@pytest.mark.asyncio
async def test_guest_ocr_records_inherit_expiry(monkeypatch):
    async def fake_process_document(self, **kwargs):
        extraction = {"transactions": [tx("2026-10-01", "NEFT ACME Corp invoice 101", credit=5000.0, ref="UTR1001")]}
        return ExtractionResponse(id=kwargs["client_request_id"], document_type="bank_statement",
                                  extraction=extraction, raw_text="...", metadata=ProcessingMetadata(pages=1))

    monkeypatch.setattr(ExtractionPipeline, "process_document", fake_process_document)
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files(bank=("s.pdf", make_pdf(1), "application/pdf")))).json()
        rid = body["reconciliation_id"]
        await wait_for(c, rid, {"X-Reconciliation-Token": body["access_token"]})
    db = MongoDBManager.get_db()
    record = await db["reconciliations"].find_one({"reconciliation_id": rid})
    job = await db["jobs"].find_one({"job_id": record["documents"]["bank_statement"]["job_id"]})
    assert job["expires_at"] == record["expires_at"]


def encrypted_pdf(user_pw: str = "u") -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "secret")
    encrypted = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw=user_pw)
    doc.close()
    return encrypted


@pytest.mark.asyncio
async def test_encrypted_pdf_without_password_asks_for_it():
    async with client() as c:
        resp = await c.post(f"{BASE}/upload", files=files(ledger=("ledger.pdf", encrypted_pdf(), "application/pdf")))
    assert resp.status_code == 400
    assert resp.json()["error"] == "Business ledger: This PDF is password-protected. Provide its password to continue."


@pytest.mark.asyncio
async def test_encrypted_pdf_with_wrong_password_is_rejected():
    async with client() as c:
        resp = await c.post(
            f"{BASE}/upload",
            files=files(bank=("bank.pdf", encrypted_pdf(), "application/pdf")),
            data={"bank_statement_password": "wrong"},
        )
    assert resp.status_code == 400
    assert resp.json()["error"] == "Bank statement: Incorrect PDF password."


@pytest.mark.asyncio
async def test_encrypted_pdf_with_password_is_accepted_and_password_not_stored():
    async with client() as c:
        resp = await c.post(
            f"{BASE}/upload",
            files=files(ledger=("ledger.pdf", encrypted_pdf("s3cret-pw"), "application/pdf")),
            data={"business_ledger_password": "s3cret-pw"},
        )
    assert resp.status_code == 202, resp.text
    db = MongoDBManager.get_db()
    record = await db["reconciliations"].find_one({"reconciliation_id": resp.json()["reconciliation_id"]})
    job = await db["jobs"].find_one({"job_id": record["documents"]["business_ledger"]["job_id"]})
    assert "s3cret-pw" not in json.dumps(record, default=str)
    assert "s3cret-pw" not in json.dumps(job, default=str)


def test_unlock_pdf_returns_unencrypted_copy():
    unlocked = ReconciliationService.unlock_pdf(encrypted_pdf(), "u")
    doc = pymupdf.open(stream=unlocked, filetype="pdf")
    assert not doc.needs_pass and not doc.is_encrypted
    assert "secret" in doc[0].get_text()
    doc.close()


# ---------------------------------------------------------------------------
# Multi-currency support
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("CA$1,234.56", "1234.56"), ("C$ 99.00", "99.00"), ("AU$1,234.56", "1234.56"), ("A$50", "50.00"),
    ("US$10", "10.00"), ("£1,234.56", "1234.56"), ("£(45.00)", "-45.00"), ("-£45.00", "-45.00"),
    ("€1.234,56", "1234.56"), ("1.234,56 €", "1234.56"), ("R$ 1.234,56", "1234.56"),
    ("CHF 1'234.50", "1234.50"), ("1 234,56 zł", "1234.56"), ("kr 1 234,56", "1234.56"), ("12,50", "12.50"),
])
def test_parse_amount_international_formats(raw, expected):
    from app.services.reconciliation_engine import parse_amount
    assert str(parse_amount(raw)[0]) == expected


@pytest.mark.parametrize("raw,code", [
    ("CAD", "CAD"), ("cad", "CAD"), ("CA$", "CAD"), ("C$", "CAD"), ("AU$", "AUD"), ("A$", "AUD"),
    ("£", "GBP"), ("GBP", "GBP"), ("€", "EUR"), ("₹", "INR"), ("R$", "BRL"), ("US$", "USD"), ("zł", "PLN"),
    ("$", None), ("¥", None), ("XYZ", None),
])
def test_normalize_currency(raw, code):
    from app.services.reconciliation_engine import normalize_currency
    assert normalize_currency(raw) == code


def test_ambiguous_dollar_does_not_conflict_with_cad():
    bank = {"currency": "$", "transactions": [tx("2026-10-01", "Sale", credit=100)]}
    ledger = {"currency": "CAD", "transactions": [tx("2026-10-01", "Sale", credit=100)]}
    result = reconcile(bank, ledger)
    assert result["summary"]["currency"] == "CAD" and result["summary"]["matched"] == 1
    assert any("shared by several currencies" in w for w in result["warnings"])


def test_cad_vs_aud_is_rejected():
    with pytest.raises(ReconciliationError, match="CAD.*AUD"):
        reconcile({"currency": "CA$", "transactions": [tx("2026-10-01", "a", credit=1)]},
                  {"currency": "A$", "transactions": [tx("2026-10-01", "a", credit=1)]})


def test_euro_decimal_comma_documents_reconcile():
    bank = {"currency": "EUR", "opening_balance": "1.000,00", "closing_balance": "2.234,56", "transactions": [
        tx("01.10.2026", "Kunde Zahlung", credit="1.234,56", balance="2.234,56"),
    ]}
    ledger = {"transactions": [tx("01.10.2026", "Kunde Zahlung", credit="1.234,56")]}
    result = reconcile(bank, ledger)
    assert result["summary"]["matched"] == 1
    assert result["balance_check"]["status"] == "passed"
    assert result["matched_transactions"][0]["amount"] == "1234.56"


def test_three_digit_groups_in_decimal_comma_document():
    # "1.234" alone is ambiguous; another value in the document shows the comma is the decimal separator.
    bank = {"transactions": [tx("2026-10-01", "a", debit="1.234"), tx("2026-10-02", "b", debit="10,50")]}
    ledger = {"transactions": [tx("2026-10-01", "a", debit=1234), tx("2026-10-02", "b", debit=10.5)]}
    assert reconcile(bank, ledger)["summary"]["matched"] == 2


@pytest.mark.parametrize("currency,amount,expected,places", [
    ("JPY", "¥12,345", "12345", 0), ("KWD", "1.234", "1.234", 3), ("GBP", "£10.5", "10.50", 2),
])
def test_currency_minor_units(currency, amount, expected, places):
    bank = {"currency": currency, "transactions": [tx("2026-10-01", "a", debit=amount)]}
    ledger = {"transactions": [tx("2026-10-01", "a", debit=amount)]}
    result = reconcile(bank, ledger)
    assert result["summary"]["amount_decimal_places"] == places
    assert result["matched_transactions"][0]["amount"] == expected


def test_kwd_three_decimals_are_not_rounded_into_false_matches():
    bank = {"currency": "KWD", "transactions": [tx("2026-10-01", "Fee", debit="1.234", ref="INV9001")]}
    ledger = {"transactions": [tx("2026-10-01", "Fee", debit="1.231", ref="INV9001")]}
    result = reconcile(bank, ledger)
    assert result["summary"]["matched"] == 0 and result["summary"]["amount_mismatches"] == 1


def test_spreadsheet_currency_detection():
    gbp = parse_spreadsheet("Date,Description,Amount\n01/10/2026,Sale,\"£1,250.00\"\n".encode(), "x.csv")
    assert gbp["currency"] == "GBP"
    cad = parse_spreadsheet("Date,Description,Amount (CAD)\n01/10/2026,Sale,1250.00\n".encode(), "x.csv")
    assert cad["currency"] == "CAD"
    aud = parse_spreadsheet("Date,Description,Amount\n01/10/2026,Sale,AU$50\n".encode(), "x.csv")
    assert aud["currency"] == "AUD"
    eur = parse_spreadsheet("Datum;Beschreibung;Betrag\n01.10.2026;Zahlung;1.234,56 €\n".encode(), "x.csv")
    assert eur["currency"] == "EUR"
    with pytest.raises(ReconciliationError, match="mixes several currencies"):
        parse_spreadsheet("Date,Description,Amount\n01/10/2026,a,CA$5\n02/10/2026,b,US$5\n".encode(), "x.csv")


@pytest.mark.asyncio
async def test_jpy_export_uses_zero_decimals():
    bank = "Date,Description,Debit,Credit\n01/10/2026,Rent,\"¥120,000\",\n"
    ledger = "Date,Description,Debit (JPY),Credit (JPY)\n01/10/2026,Rent,120000,\n"
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files(
            bank=("b.csv", bank.encode(), "text/csv"), ledger=("l.csv", ledger.encode(), "text/csv")))).json()
        rid, headers = body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]}
        final = await wait_for(c, rid, headers)
        assert final["result"]["summary"]["currency"] == "JPY"
        assert final["result"]["matched_transactions"][0]["amount"] == "120000"
        wb = openpyxl.load_workbook(io.BytesIO((await c.get(f"{BASE}/{rid}/export", headers=headers)).content))
        assert wb["Matched"].cell(row=2, column=9).number_format == "#,##0"


# ---------------------------------------------------------------------------
# OCR amount verification against raw text
# ---------------------------------------------------------------------------

EU_RAW_TEXT = """Kontoauszug  Währung: EUR
Anfangssaldo 1.234,56
01.10.2026 Gehalt ACME GmbH          2.500,00     3.734,56
03.10.2026 Miete Oktober    1.234,56              2.500,00
05.10.2026 Kontoführung        12,50              2.487,50
Endsaldo 2.487,50
"""


def test_misread_european_amounts_are_corrected_from_raw_text():
    from app.services.reconciliation_engine import verify_amounts_against_raw_text
    data = {"_raw_text": EU_RAW_TEXT, "opening_balance": 1.23, "closing_balance": 2487.5, "transactions": [
        tx("2026-10-01", "Gehalt ACME", credit=2.5, balance=3734.56),
        tx("2026-10-03", "Miete Oktober", debit=123456.0, balance=2500.0),
        tx("2026-10-05", "Kontoführung", debit=12.5, balance=2487.5),
    ]}
    fixed, corrections, warnings = verify_amounts_against_raw_text(data, "bank")
    assert {(c["transaction_id"], c["field"], c["extracted"], c["corrected"]) for c in corrections} == {
        (None, "opening_balance", "1.23", "1234.56"),
        ("B1", "credit", "2.50", "2500.00"),
        ("B2", "debit", "123456.00", "1234.56"),
    }
    assert data["transactions"][0]["credit"] == 2.5  # original extraction untouched
    assert any("misread" in w for w in warnings)


def test_correct_extraction_and_dot_decimal_documents_are_left_alone():
    from app.services.reconciliation_engine import verify_amounts_against_raw_text
    data = {"_raw_text": EU_RAW_TEXT, "opening_balance": 1234.56, "transactions": [
        tx("2026-10-01", "Gehalt", credit=2500.0), tx("2026-10-05", "Gebühr", debit=12.5)]}
    assert verify_amounts_against_raw_text(data, "bank")[1] == []

    us_raw = "Opening 1,234.56\n10/01/2026 Payroll 2,500.00 3,734.56\n10/03/2026 Rent 1,234.56 2,500.00\n"
    data = {"_raw_text": us_raw, "transactions": [tx("2026-10-01", "Payroll", credit=2500.0)]}
    assert verify_amounts_against_raw_text(data, "bank")[1] == []


def test_ambiguous_misreading_is_not_guessed():
    from app.services.reconciliation_engine import verify_amounts_against_raw_text
    # 1.23 could come from "1.234,56" or "1.230,00": two candidates, so nothing is changed.
    raw = "Saldo 1.234,56\nZahlung 1.230,00\nGebühr 5,50\nGebühr 7,25\n"
    data = {"_raw_text": raw, "transactions": [tx("2026-10-01", "x", debit=1.23)]}
    fixed, corrections, _ = verify_amounts_against_raw_text(data, "bank")
    assert corrections == [] and fixed["transactions"][0]["debit"] == 1.23


@pytest.mark.asyncio
async def test_scanned_european_statement_reconciles_after_correction(monkeypatch):
    async def fake_process_document(self, **kwargs):
        extraction = {
            "currency": "EUR", "opening_balance": 1.23, "closing_balance": 2487.5,
            "transactions": [
                tx("2026-10-01", "Gehalt ACME GmbH", credit=2.5, balance=3734.56),
                tx("2026-10-03", "Miete Oktober", debit=123456.0, balance=2500.0),
                tx("2026-10-05", "Kontoführung", debit=12.5, balance=2487.5),
            ],
        }
        return ExtractionResponse(id=kwargs["client_request_id"], document_type="bank_statement",
                                  extraction=extraction, raw_text=EU_RAW_TEXT, metadata=ProcessingMetadata(pages=1))

    monkeypatch.setattr(ExtractionPipeline, "process_document", fake_process_document)
    ledger = (
        "Datum;Buchungstext;Soll;Haben;Saldo\n"
        "01.10.2026;Anfangssaldo;;;1.234,56\n"
        "01.10.2026;Gehalt ACME GmbH;;2.500,00;3.734,56\n"
        "03.10.2026;Miete Oktober;1.234,56;;2.500,00\n"
        "05.10.2026;Kontoführung;12,50;;2.487,50\n"
    )
    async with client() as c:
        body = (await c.post(f"{BASE}/upload", files=files(
            bank=("kontoauszug.pdf", make_pdf(1), "application/pdf"),
            ledger=("buchungen.csv", ledger.encode(), "text/csv")))).json()
        final = await wait_for(c, body["reconciliation_id"], {"X-Reconciliation-Token": body["access_token"]})
    result = final["result"]
    assert result["summary"]["matched"] == 3
    assert result["summary"]["unmatched_bank"] == 0 and result["summary"]["amount_mismatches"] == 0
    assert result["summary"]["amount_corrections"] == 3
    assert result["balance_check"]["status"] == "passed"
    assert result["balance_check"]["book_to_bank"]["status"] == "passed"
    assert len(result["amount_corrections"]) == 3
