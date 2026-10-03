"""
Tests for Document Export Service and Export API Endpoints (Excel, CSV, PDF).
"""

import io
import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient

from app.main import app
from app.services.export_service import ExportService
from app.services.result_cache import ResultCache


@pytest.fixture
def sample_bank_extraction():
    return {
        "bank_name": "State Bank of India",
        "account_holder": "Mr. NAVNIT RAI",
        "account_number_masked": "XXXX-123456",
        "currency": "INR",
        "statement_period": "2026-09-01 to 2026-09-30",
        "opening_balance": 10000.00,
        "closing_balance": 34500.00,
        "transactions": [
            {
                "date": "2026-09-01",
                "description": "UPI PAYMENT",
                "reference": "UPI123456",
                "debit": 500.00,
                "credit": None,
                "balance": 9500.00,
            },
            {
                "date": "2026-09-02",
                "description": "SALARY CREDIT",
                "reference": "SAL9988",
                "debit": None,
                "credit": 25000.00,
                "balance": 34500.00,
            },
        ],
    }


@pytest.fixture
def sample_invoice_extraction():
    return {
        "invoice_number": "INV-2026-8800",
        "invoice_date": "2026-09-10",
        "due_date": "2026-10-10",
        "currency": "USD",
        "supplier": {"name": "Cloud Infra Inc.", "address": "123 Tech Blvd", "tax_id": "US-889900"},
        "customer": {"name": "Alpha Corp", "address": "456 Market St", "tax_id": "US-112233"},
        "subtotal": 5000.00,
        "tax": 450.00,
        "total": 5450.00,
        "line_items": [
            {"description": "Server Hosting", "quantity": 2, "unit_price": 2000.00, "tax_rate": "9%", "amount": 4000.00},
            {"description": "Database Backup", "quantity": 1, "unit_price": 1000.00, "tax_rate": "9%", "amount": 1000.00},
        ],
    }


def test_analytics_categorization_and_subscriptions(sample_bank_extraction):
    from app.services.analytics_service import AnalyticsService
    analytics = AnalyticsService.analyze("bank_statement", sample_bank_extraction)

    assert analytics.total_inflow == 25000.00
    assert analytics.total_outflow == 500.00
    assert analytics.net_savings == 24500.00
    assert analytics.transaction_count == 2
    assert len(analytics.ai_insights) > 0

    # Test category matching
    assert AnalyticsService.categorize_description("NETFLIX.COM PAYMENT") == "Entertainment & Subscriptions"
    assert AnalyticsService.categorize_description("STARBUCKS COFFEE") == "Food & Dining"
    assert AnalyticsService.categorize_description("MONTHLY SALARY TRANSFER") == "Salary & Income"
    assert AnalyticsService.categorize_description("UBER TRIP") == "Travel & Commute"


def test_generate_excel_bank_statement(sample_bank_extraction):
    file_bytes = ExportService.generate_excel(
        doc_id="test_doc_01",
        document_type="bank_statement",
        extraction=sample_bank_extraction,
    )
    assert isinstance(file_bytes, bytes)
    assert len(file_bytes) > 0
    assert file_bytes[:2] == b"PK"

    # Verify workbook structure and no diagonal column drift
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    assert "Executive Dashboard" in wb.sheetnames
    assert "Itemized Ledger" in wb.sheetnames

    ws_ledger = wb["Itemized Ledger"]
    # Row 4 is header, Row 5 is Tx 1, Row 6 is Tx 2, Row 7 is Summary
    assert ws_ledger.cell(row=5, column=1).value == "2026-09-01"
    assert ws_ledger.cell(row=5, column=3).value == "UPI PAYMENT"
    assert ws_ledger.cell(row=5, column=5).value == 500.00

    assert ws_ledger.cell(row=6, column=1).value == "2026-09-02"
    assert ws_ledger.cell(row=6, column=3).value == "SALARY CREDIT"
    assert ws_ledger.cell(row=6, column=6).value == 25000.00


def test_generate_excel_invoice(sample_invoice_extraction):
    file_bytes = ExportService.generate_excel(
        doc_id="test_inv_01",
        document_type="invoice",
        extraction=sample_invoice_extraction,
    )
    assert isinstance(file_bytes, bytes)
    assert file_bytes[:2] == b"PK"

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws_ledger = wb["Itemized Ledger"]
    # Verify line items on separate rows
    assert ws_ledger.cell(row=5, column=1).value == "Server Hosting"
    assert ws_ledger.cell(row=5, column=3).value == 2
    assert ws_ledger.cell(row=5, column=6).value == 4000.00

    assert ws_ledger.cell(row=6, column=1).value == "Database Backup"
    assert ws_ledger.cell(row=6, column=3).value == 1
    assert ws_ledger.cell(row=6, column=6).value == 1000.00


def test_generate_csv(sample_bank_extraction):
    csv_bytes = ExportService.generate_csv(
        doc_id="test_doc_01",
        document_type="bank_statement",
        extraction=sample_bank_extraction,
    )
    assert isinstance(csv_bytes, bytes)
    csv_text = csv_bytes.decode("utf-8-sig")
    assert "State Bank of India" in csv_text
    assert "Mr. NAVNIT RAI" in csv_text
    assert "UPI PAYMENT" in csv_text


def test_generate_pdf(sample_bank_extraction):
    pdf_bytes = ExportService.generate_pdf(
        doc_id="test_doc_01",
        document_type="bank_statement",
        extraction=sample_bank_extraction,
    )
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 0
    assert pdf_bytes.startswith(b"%PDF-")


def test_generate_pdf_invoice(sample_invoice_extraction):
    pdf_bytes = ExportService.generate_pdf(
        doc_id="test_inv_01",
        document_type="invoice",
        extraction=sample_invoice_extraction,
    )
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_export_download_by_id_e2e(sample_bank_extraction):
    from app.core.security import hash_key
    doc_id = "cache_test_id_7788"
    auth_headers = {"X-API-Key": "ocr_dev_key_secret_2026"}
    owner_hash = hash_key("ocr_dev_key_secret_2026")

    # Seed cache
    ResultCache.set(
        doc_id,
        {
            "id": doc_id,
            "document_type": "bank_statement",
            "extraction": sample_bank_extraction,
            "raw_text": "Sample text",
        },
        owner_hash=owner_hash,
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Download Excel
        res_xlsx = await client.get(f"/api/v1/export/download/{doc_id}?format=xlsx", headers=auth_headers)
        assert res_xlsx.status_code == 200
        assert "application/vnd.openxmlformats" in res_xlsx.headers["content-type"]
        assert "bank_statement_cache_test_id_7788.xlsx" in res_xlsx.headers["content-disposition"]
        assert res_xlsx.content[:2] == b"PK"

        # 2. Download CSV
        res_csv = await client.get(f"/api/v1/export/download/{doc_id}?format=csv", headers=auth_headers)
        assert res_csv.status_code == 200
        assert "text/csv" in res_csv.headers["content-type"]
        assert "bank_statement_cache_test_id_7788.csv" in res_csv.headers["content-disposition"]

        # 3. Download PDF
        res_pdf = await client.get(f"/api/v1/export/download/{doc_id}?format=pdf", headers=auth_headers)
        assert res_pdf.status_code == 200
        assert "application/pdf" in res_pdf.headers["content-type"]
        assert res_pdf.content.startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_export_download_not_found():
    auth_headers = {"X-API-Key": "ocr_dev_key_secret_2026"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get("/api/v1/export/download/nonexistent_id_9999?format=xlsx", headers=auth_headers)
        assert res.status_code == 404
        assert "not found or has expired" in res.json()["error"]


@pytest.mark.asyncio
async def test_export_direct_generate_e2e(sample_bank_extraction):
    auth_headers = {"X-API-Key": "ocr_dev_key_secret_2026"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        payload = {
            "id": "direct_req_123",
            "document_type": "bank_statement",
            "extraction": sample_bank_extraction,
            "raw_text": "Sample direct text",
        }
        res = await client.post("/api/v1/export/generate?format=xlsx", json=payload, headers=auth_headers)
        assert res.status_code == 200
        assert res.content[:2] == b"PK"
        assert "bank_statement_direct_req_123.xlsx" in res.headers["content-disposition"]


def test_excel_export_with_string_formatted_numbers():
    """Verify that string-formatted amounts ($1,200.50, 5,000) are parsed as real numbers in Excel."""
    extraction_with_strings = {
        "bank_name": "Chase Bank",
        "account_holder": "Jane Doe",
        "opening_balance": "10,000.00",
        "closing_balance": "$15,250.50",
        "transactions": [
            {
                "date": "2026-05-01",
                "description": "=cmd|' /C calc'!A0",
                "reference": "REF001",
                "debit": "1,500.00",
                "credit": None,
                "balance": "8,500.00",
            },
            {
                "date": "2026-05-05",
                "description": "Salary Deposit",
                "reference": "SAL002",
                "debit": None,
                "credit": "6,750.50",
                "balance": "15,250.50",
            },
        ],
    }

    excel_bytes = ExportService.generate_excel(
        doc_id="str_test_1",
        document_type="bank_statement",
        extraction=extraction_with_strings,
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Itemized Ledger"]

    # Verify formula injection sanitization
    assert str(ws.cell(row=5, column=3).value).startswith("'=")

    # Verify numbers are numeric and not strings
    assert ws.cell(row=5, column=5).value == 1500.00
    assert isinstance(ws.cell(row=5, column=5).value, (int, float))
    assert ws.cell(row=6, column=6).value == 6750.50
    assert isinstance(ws.cell(row=6, column=6).value, (int, float))

    # Verify summary total row
    assert ws.cell(row=7, column=5).value == 1500.00
    assert ws.cell(row=7, column=6).value == 6750.50
    assert ws.cell(row=7, column=7).value == 15250.50


def test_consolidation_excel_with_invoices_and_multi_format_dates():
    """Verify that consolidation Excel supports invoices, diverse date formats, and summary row."""
    from app.services.consolidation_service import ConsolidationService

    extractions = [
        {
            "currency": "USD",
            "transactions": [
                {
                    "date": "15/09/2026",  # DD/MM/YYYY
                    "description": "Software Subscription",
                    "reference": "SUB100",
                    "debit": "250.00",
                    "credit": None,
                    "balance": "5000.00",
                }
            ],
        },
        {
            "currency": "USD",
            "invoice_number": "INV-9900",
            "invoice_date": "20-Oct-2026",  # DD-Mon-YYYY
            "line_items": [
                {
                    "description": "Consulting Services",
                    "quantity": 10,
                    "unit_price": "150.00",
                    "amount": "1500.00",
                }
            ],
        },
    ]

    consolidation = ConsolidationService.consolidate(extractions)
    assert consolidation.statement_count == 2
    assert consolidation.total_transactions == 2
    assert len(consolidation.monthly_trends) == 2

    # Check months are 2026-09 and 2026-10
    months = [m.month for m in consolidation.monthly_trends]
    assert "2026-09" in months
    assert "2026-10" in months

    # Generate workbook and verify Master Ledger has both items
    wb_bytes = ConsolidationService.generate_consolidated_excel(consolidation, extractions)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    assert "Annual Consolidation" in wb.sheetnames
    assert "Master Ledger" in wb.sheetnames

    ws_ledg = wb["Master Ledger"]
    assert ws_ledg.cell(row=5, column=3).value == "Software Subscription"
    assert ws_ledg.cell(row=5, column=5).value == 250.00

    assert ws_ledg.cell(row=6, column=3).value == "Consulting Services"
    assert ws_ledg.cell(row=6, column=5).value == 1500.00

    # Summary row in Master Ledger
    assert ws_ledg.cell(row=7, column=1).value == "TOTAL SUMMARY"
    assert ws_ledg.cell(row=7, column=5).value == 1750.00

